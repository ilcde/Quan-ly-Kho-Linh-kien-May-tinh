from pathlib import Path
from copy import deepcopy
from zipfile import ZipFile, ZIP_DEFLATED
import hashlib, json
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image, ImageDraw, ImageFont

ROOT=Path(__file__).resolve().parent.parent
REF=Path('C:/Users/ilcde/Music/Bao cao do an java.docx')
OUT=ROOT/'docs'; OUT.mkdir(exist_ok=True)
original=hashlib.sha256(REF.read_bytes()).hexdigest()
(OUT/'reference-evidence.json').write_text(json.dumps({'source':str(REF),'sha256':original,'pages':21,'section_count':1},indent=2),encoding='utf-8')

def erd():
    im=Image.new('RGB',(1600,1100),'white');d=ImageDraw.Draw(im)
    font=ImageFont.truetype('C:/Windows/Fonts/arial.ttf',25)
    bold=ImageFont.truetype('C:/Windows/Fonts/arialbd.ttf',28)
    boxes={
      'NhaCungCap':(30,30,['PK MaNCC','TenNCC, DienThoai, DiaChi']),
      'PhieuNhap':(30,345,['PK MaPN','FK MaNCC','NgayTao, NguoiLap, GhiChu']),
      'ChiTietNhap':(30,690,['PK/FK MaPN','PK/FK MaLK','SoLuong, DonGia']),
      'LoaiLinhKien':(590,30,['PK MaLoai','TenLoai']),
      'LinhKien':(590,380,['PK MaLK','FK MaLoai','TenLK, Hang, DonVi','GiaBan, TonToiThieu']),
      'PhieuXuat':(1150,345,['PK MaPX','NguoiNhan, NgayTao','NguoiLap, GhiChu']),
      'ChiTietXuat':(1150,690,['PK/FK MaPX','PK/FK MaLK','SoLuong, DonGia'])}
    def link(points,a,b):
        d.line(points,fill='#333333',width=4);d.text(a,'1',font=bold,fill='black');d.text(b,'N',font=bold,fill='black')
    link([(235,190),(235,345)],(248,205),(248,300))
    link([(235,540),(235,690)],(248,550),(248,650))
    link([(795,190),(795,380)],(808,205),(808,338))
    link([(1355,540),(1355,690)],(1368,550),(1368,650))
    link([(590,510),(500,510),(500,800),(440,800)],(550,470),(455,760))
    link([(1000,510),(1080,510),(1080,800),(1150,800)],(1020,470),(1100,760))
    for title,(x,y,fields) in boxes.items():
        h=70+len(fields)*37;d.rectangle((x,y,x+410,y+h),fill='white',outline='#4174a8',width=3);d.rectangle((x,y,x+410,y+52),fill='#dce6f1');d.text((x+12,y+10),title,font=bold,fill='black')
        for i,text in enumerate(fields):d.text((x+12,y+67+i*37),text,font=font,fill='black')
    d.text((30,1000),'PK: khoa chinh     FK: khoa ngoai     1—N: quan he mot nhieu',font=font,fill='black')
    im.save(OUT/'images/erd.png')
erd()

schema=[
('LoaiLinhKien',[
('MaLoai','nvarchar(20)','Khóa chính; mã loại duy nhất, không rỗng.'),('TenLoai','nvarchar(100)','Bắt buộc, không rỗng và duy nhất.')]),
('NhaCungCap',[
('MaNCC','nvarchar(20)','Khóa chính; không rỗng.'),('TenNCC','nvarchar(150)','Tên nhà cung cấp, bắt buộc.'),('DienThoai','nvarchar(20)','Số điện thoại; mặc định chuỗi rỗng.'),('DiaChi','nvarchar(250)','Địa chỉ; mặc định chuỗi rỗng.')]),
('LinhKien',[
('MaLK','nvarchar(20)','Khóa chính; không rỗng.'),('TenLK','nvarchar(150)','Tên linh kiện; bắt buộc, không rỗng.'),('MaLoai','nvarchar(20)','Khóa ngoại tới LoaiLinhKien.MaLoai.'),('Hang','nvarchar(80)','Hãng sản xuất; bắt buộc ở giao diện.'),('DonVi','nvarchar(20)','Đơn vị tính; bắt buộc ở giao diện.'),('GiaBan','decimal(18,2)','Giá bán tham khảo; CHECK >= 0.'),('TonToiThieu','int','Mức cảnh báo; CHECK >= 0; mặc định 5.')]),
('PhieuNhap',[
('MaPN','int IDENTITY','Khóa chính, tự tăng.'),('NgayTao','datetime2','Mặc định SYSDATETIME tại máy chủ.'),('MaNCC','nvarchar(20)','Khóa ngoại tới NhaCungCap.MaNCC.'),('NguoiLap','nvarchar(100)','Người lập phiếu; không rỗng.'),('GhiChu','nvarchar(250)','Ghi chú; mặc định chuỗi rỗng.')]),
('ChiTietNhap',[
('MaPN','int','Khóa chính ghép; FK tới PhieuNhap.MaPN.'),('MaLK','nvarchar(20)','Khóa chính ghép; FK tới LinhKien.MaLK.'),('SoLuong','int','Số lượng nhập; CHECK > 0.'),('DonGia','decimal(18,2)','Đơn giá nhập; CHECK >= 0.')]),
('PhieuXuat',[
('MaPX','int IDENTITY','Khóa chính, tự tăng.'),('NgayTao','datetime2','Mặc định SYSDATETIME tại máy chủ.'),('NguoiNhan','nvarchar(150)','Tên người nhận; không rỗng.'),('NguoiLap','nvarchar(100)','Người lập phiếu; không rỗng.'),('GhiChu','nvarchar(250)','Ghi chú; mặc định chuỗi rỗng.')]),
('ChiTietXuat',[
('MaPX','int','Khóa chính ghép; FK tới PhieuXuat.MaPX.'),('MaLK','nvarchar(20)','Khóa chính ghép; FK tới LinhKien.MaLK.'),('SoLuong','int','Số lượng xuất; CHECK > 0.'),('DonGia','decimal(18,2)','Đơn giá xuất; CHECK >= 0.')])]

def newdoc():
    doc=Document(REF)
    body=doc._element.body
    for child in list(body):
        if child.tag!=qn('w:sectPr'):body.remove(child)
    for name in ['Normal','Normal (Web)','Heading 1','Heading 2','Heading 3','Title']:
        if name not in doc.styles:doc.styles.add_style(name,WD_STYLE_TYPE.PARAGRAPH)
        s=doc.styles[name];s.font.name='Times New Roman';s.font.size=Pt(14);s.font.color.rgb=RGBColor(0,0,0)
        s.paragraph_format.space_after=Pt(8);s.paragraph_format.line_spacing=1.15
    for name in ['Heading 1','Heading 2','Heading 3']:
        doc.styles[name].font.bold=True;doc.styles[name].paragraph_format.keep_with_next=True
        pr=doc.styles[name].element.get_or_add_pPr();outline=OxmlElement('w:outlineLvl');outline.set(qn('w:val'),str(int(name[-1])-1));pr.append(outline)
    return doc
pending_pages=set()
def apply_page(doc,x):
    if id(doc) in pending_pages:x.paragraph_format.page_break_before=True;pending_pages.remove(id(doc))
    return x
def p(doc,text='',bold=False):
    x=apply_page(doc,doc.add_paragraph(style='Normal (Web)'));r=x.add_run(text);r.bold=bold;return x
def h(doc,text,level=1):apply_page(doc,doc.add_heading(text,level))
def page(doc):pending_pages.add(id(doc))
def pic(doc,path,caption):
    x=doc.add_paragraph();x.alignment=WD_ALIGN_PARAGRAPH.CENTER;x.add_run().add_picture(str(path),width=Inches(6.45));
    c=p(doc,caption);c.alignment=WD_ALIGN_PARAGRAPH.CENTER;c.runs[0].italic=True;c.runs[0].font.size=Pt(12)
def table(doc,headers,rows):
    t=doc.add_table(rows=1,cols=len(headers));t.style='Light Grid Accent 1';t.autofit=False
    widths=[1.25,1.55,3.65] if len(headers)==3 else [1.2,5.25]
    for c,w in zip(t.columns,widths):c.width=Inches(w)
    for i,v in enumerate(headers):t.rows[0].cells[i].text=v
    pr=t.rows[0]._tr.get_or_add_trPr();pr.append(OxmlElement('w:tblHeader'))
    for row in rows:
        for c,v in zip(t.add_row().cells,row):c.text=str(v)
    for row in t.rows:
        trpr=row._tr.get_or_add_trPr();trpr.append(OxmlElement('w:cantSplit'))
        for cell,w in zip(row.cells,widths):
            cell.width=Inches(w)
            for para in cell.paragraphs:
                para.paragraph_format.space_after=Pt(4);para.paragraph_format.space_before=Pt(4);para.paragraph_format.line_spacing=1
                for r in para.runs:r.font.name='Times New Roman';r.font.size=Pt(12)
    return t
def toc(doc):
    p(doc,'MỤC LỤC',True)
    x=p(doc);f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'TOC \\o "1-2" \\h \\z \\u');x._p.append(f)
def save(doc,name):
    # Preserve source package parts untouched outside declared content/style/relationship edits.
    path=OUT/name;temp=OUT/('_'+name);doc.save(temp)
    with ZipFile(REF) as src,ZipFile(temp) as generated,ZipFile(path,'w',ZIP_DEFLATED) as out:
        changed={'word/document.xml','word/_rels/document.xml.rels','word/styles.xml','word/settings.xml','[Content_Types].xml'}
        srcnames=set(src.namelist())
        for item in src.infolist():out.writestr(item,generated.read(item.filename) if item.filename in changed else src.read(item.filename))
        for item in generated.infolist():
            if item.filename not in srcnames:out.writestr(item,generated.read(item.filename))
    temp.unlink()
    assert hashlib.sha256(REF.read_bytes()).hexdigest()==original
    return path
def cover(doc,title):
    p(doc,'ĐỒ ÁN MÔN HỌC JAVA',True).alignment=WD_ALIGN_PARAGRAPH.CENTER
    p(doc);p(doc)
    x=doc.add_paragraph(title,style='Title');x.alignment=WD_ALIGN_PARAGRAPH.CENTER;x.runs[0].bold=True;x.runs[0].font.size=Pt(20)
    p(doc);p(doc,'Sinh viên thực hiện: Nguyễn Đức Tuệ');p(doc,'Lớp: TH30.18');p(doc,'Mã sinh viên: ........................................');p(doc,'Giảng viên hướng dẫn: ........................................')
    p(doc);p(doc,'Năm 2026').alignment=WD_ALIGN_PARAGRAPH.CENTER
def database(doc,prefix='2.1.2'):
    h(doc,prefix+' Thiết kế cơ sở dữ liệu',2)
    p(doc,'Cơ sở dữ liệu KhoLinhKienJava gồm bảy bảng. Loại linh kiện được tách khỏi thông tin linh kiện; mỗi phiếu gồm bảng đầu phiếu và bảng chi tiết. Cách tổ chức này tránh lặp lại ngày tạo, đối tác và người lập ở từng dòng hàng.')
    pic(doc,OUT/'images/erd.png','Sơ đồ quan hệ của bảy bảng trong cơ sở dữ liệu')
    p(doc,'Mỗi loại có nhiều linh kiện; mỗi nhà cung cấp có nhiều phiếu nhập. Một phiếu có nhiều dòng chi tiết, mỗi dòng liên kết tới một linh kiện. Một linh kiện có thể xuất hiện trên nhiều phiếu. Khóa chính ghép ở chi tiết bảo đảm một linh kiện chỉ xuất hiện một lần trong một phiếu.')
    page(doc)
    for i,(name,rows) in enumerate(schema):
        h(doc,'Bảng '+name,3);table(doc,['Tên trường','Kiểu dữ liệu','Mô tả và ràng buộc'],rows)
        if i in [1,2,4]:page(doc)
    page(doc)
    h(doc,'Ràng buộc và cách tính tồn kho',2)
    p(doc,'Tất cả cột đều NOT NULL. Các trường điện thoại, địa chỉ và ghi chú có thể nhận chuỗi rỗng. Mã và tên cốt lõi phải có nội dung. Khóa ngoại sử dụng hành vi mặc định NO ACTION: không xóa loại, linh kiện hoặc nhà cung cấp khi dữ liệu còn được tham chiếu.')
    p(doc,'Giá dùng decimal(18,2) để tránh sai số của số thực nhị phân. Số lượng mỗi dòng là số nguyên dương. Tổng số lượng được cộng dưới kiểu bigint trong view để hạn chế tràn khi có nhiều phiếu.')
    p(doc,'View v_TonKho tính tồn của từng linh kiện bằng tổng ChiTietNhap.SoLuong trừ tổng ChiTietXuat.SoLuong. Các tổng được gom riêng trước khi JOIN, tránh nhân bản số lượng khi một linh kiện có nhiều dòng nhập và nhiều dòng xuất. Linh kiện chưa phát sinh phiếu có tồn bằng 0.')
    p(doc,'Quy tắc không xuất quá tồn được thực hiện ở InventoryService trong một giao dịch. Service khóa bản ghi LinhKien bằng UPDLOCK và HOLDLOCK theo thứ tự mã, kiểm tra tồn, ghi đầu phiếu và chi tiết, sau đó commit. Nếu một bước lỗi, toàn bộ giao dịch rollback. Khóa CHECK của từng bảng không tự kiểm soát được tổng nhập trừ tổng xuất; do đó mọi thao tác ghi sổ của ứng dụng phải đi qua service này.')
    p(doc,'Chỉ mục được tạo cho MaLoai trong LinhKien và MaLK ở hai bảng chi tiết. Các chỉ mục hỗ trợ truy vấn quan hệ và tổng hợp tồn kho. Phiếu đã ghi sổ không được sửa hoặc xóa trực tiếp trong giao diện phiên bản này.')

doc=newdoc();cover(doc,'Phần mềm Quản lý Kho Linh kiện Máy tính');page(doc);toc(doc);page(doc)
h(doc,'Chương 1 Tổng quan về phần mềm Java')
h(doc,'1.1 Giới thiệu đề tài',2)
p(doc,'Phần mềm Quản lý Kho Linh kiện Máy tính phục vụ việc theo dõi danh mục linh kiện, nhà cung cấp và hoạt động nhập xuất tại một kho nhỏ. Đối tượng quản lý gồm CPU, RAM, SSD, card đồ họa và bo mạch chủ. Người quản lý cần biết linh kiện nào đang có trong kho, số lượng còn lại và những mặt hàng cần nhập thêm.')
p(doc,'Nếu quản lý bằng ghi chép rời rạc, số liệu nhập và xuất có thể không thống nhất. Phần mềm tập trung dữ liệu vào SQL Server, cung cấp các form nhập liệu và lưu lịch sử phiếu. Mục tiêu của đồ án là xây dựng một chương trình Java có thể thêm, sửa, xóa, tìm kiếm dữ liệu, đồng thời kiểm tra những quy tắc cơ bản của nghiệp vụ kho.')
h(doc,'1.2 Các kiến thức sử dụng trong bài',2)
h(doc,'1.2.1 Cơ sở dữ liệu',3)
p(doc,'SQL Server lưu dữ liệu dưới dạng các bảng liên kết bằng khóa chính và khóa ngoại. SELECT dùng truy vấn; INSERT thêm bản ghi; UPDATE sửa dữ liệu; DELETE xóa bản ghi hợp lệ. JOIN kết hợp thông tin các bảng. GROUP BY và SUM hỗ trợ tính tồn kho. Giao dịch bảo đảm nhiều câu lệnh của một phiếu cùng thành công hoặc cùng bị hủy.')
h(doc,'1.2.2 Ngôn ngữ lập trình',3)
p(doc,'Java 21 được dùng để xây dựng chương trình. Swing cung cấp JFrame, JPanel, JDialog, JTable, JTextField và JComboBox; AWT hỗ trợ bố cục, màu sắc và phông chữ. JDBC kết nối Java với SQL Server thông qua Microsoft JDBC Driver. PreparedStatement truyền tham số truy vấn; try-with-resources đóng kết nối và ResultSet sau khi sử dụng.')
page(doc)
h(doc,'Chương 2 Phân tích và thiết kế')
h(doc,'2.1 Phân tích và thiết kế hệ thống',2)
h(doc,'2.1.1 Phân tích chương trình',3)
p(doc,'Người sử dụng là nhân viên hoặc người quản lý kho đang đăng nhập Windows. Phiên bản đồ án chạy tại một máy và một kho. Người dùng truy cập sáu màn hình qua thanh điều hướng bên trái. Ứng dụng chưa có phân quyền theo tài khoản riêng.')
table(doc,['Chức năng','Dữ liệu và thao tác'],[
('Loại linh kiện','Thêm, sửa, xóa và tìm loại; tên loại duy nhất.'),('Nhà cung cấp','Quản lý mã, tên, điện thoại, địa chỉ.'),('Linh kiện','Quản lý loại, hãng, đơn vị, giá bán, ngưỡng tồn.'),('Phiếu nhập','Chọn nhà cung cấp; nhập nhiều dòng hàng; ghi sổ.'),('Phiếu xuất','Nhập người nhận; kiểm tra tồn trước khi ghi sổ.'),('Tồn kho','Xem tổng nhập trừ tổng xuất; cảnh báo; xuất CSV.')])
p(doc,'Luồng nghiệp vụ bắt đầu từ việc tạo loại linh kiện, nhà cung cấp và danh mục linh kiện. Nhân viên lập phiếu nhập để ghi nhận hàng về, sau đó lập phiếu xuất khi giao hàng. Khi cần tra cứu, người dùng tìm theo mã hoặc tên/đối tác, chọn một phiếu để xem chi tiết hoặc xuất bảng hiện tại thành CSV.')
p(doc,'Yêu cầu dữ liệu gồm mã duy nhất, trường bắt buộc không rỗng, số lượng dương và giá không âm. Không được xuất vượt tồn. Khi lỗi kết nối hoặc vi phạm ràng buộc, giao diện thông báo lỗi và không báo thành công. Dữ liệu có lịch sử liên quan không được xóa để tránh mất dấu vết nghiệp vụ.')
page(doc);database(doc)
page(doc);h(doc,'2.2 Xây dựng các chức năng',2)
p(doc,'Mã nguồn gồm ba lớp chính. Main xây dựng cửa sổ và các form, nhận thao tác người dùng và hiển thị kết quả. Db đọc cấu hình, mở kết nối, chạy truy vấn có tham số và đóng tài nguyên. InventoryService tập trung nghiệp vụ nhập xuất và quản lý giao dịch. Việc tách nghiệp vụ kho khỏi sự kiện giao diện giúp có thể kiểm thử mà không thao tác chuột.')
h(doc,'Quản lý dữ liệu danh mục',3)
p(doc,'Khi thêm, chương trình kiểm tra độ dài, trường bắt buộc và kiểu số trước khi INSERT. Khi sửa, mã được giữ cố định và UPDATE các thông tin còn lại. Khi xóa, người dùng phải xác nhận; SQL Server kiểm tra các khóa ngoại. Tìm kiếm dùng LIKE trên mã và tên, truyền giá trị bằng PreparedStatement để dữ liệu nhập không trở thành lệnh SQL.')
h(doc,'Lập phiếu nhập và phiếu xuất',3)
p(doc,'Form phiếu có thông tin đối tác, người lập, ghi chú và bảng chi tiết. Mỗi dòng chọn một linh kiện, số lượng và đơn giá. Người dùng có thể bỏ dòng để nhập lại trước khi ghi sổ. Chương trình tính tổng tiền bằng BigDecimal. Sau xác nhận, form tạm khóa thao tác để hạn chế gửi trùng, rồi gọi service trong luồng nền.')
h(doc,'Hiển thị và xử lý nền',3)
p(doc,'SwingWorker thực hiện truy vấn ngoài luồng giao diện, giúp cửa sổ vẫn phản hồi trong lúc chờ SQL Server. Chỉ kết quả của lần tải mới nhất được đưa vào bảng. Các JTable cho phép sắp xếp; bảng tồn kho đánh dấu linh kiện có tồn nhỏ hơn hoặc bằng ngưỡng. CSV dùng UTF-8 có BOM để thuận tiện mở bằng Excel và xử lý dấu tiếng Việt.')
page(doc);h(doc,'Chương 3 Phần mềm thử nghiệm')
h(doc,'3.1 Môi trường và cách chạy',2)
p(doc,'Môi trường thử nghiệm sử dụng Windows, JDK 21.0.8, SQL Server 2025 Express và Microsoft JDBC Driver 13.6.0 jre11. Cơ sở dữ liệu KhoLinhKienJava đã được tạo từ script của dự án. SQL Server chỉ lắng nghe địa chỉ 127.0.0.1 ở cổng 1433; JDBC dùng localhost và lựa chọn IPv4. Kết nối dùng tài khoản Windows hiện tại.')
p(doc,'Để khởi động trên máy đã cấu hình, mở Run.bat trong thư mục dự án. Tệp run.ps1 biên dịch mã, tạo JAR và chạy lớp Main với thư viện trong lib. Có thể kiểm tra kết nối bằng lệnh ./run.ps1 --check-db. Lệnh ./test.ps1 chạy bộ kiểm thử tích hợp.')
p(doc,'Trên máy mới, cần cài JDK và SQL Server, chạy script schema và dữ liệu minh họa, sau đó sao chép config.example.properties thành config.properties và sửa địa chỉ kết nối. Script tạo dữ liệu không được chạy lại trên cơ sở dữ liệu đang sử dụng. Mật khẩu SQL, nếu sử dụng, được truyền bằng biến môi trường và không đưa lên GitHub.')
p(doc,'Dữ liệu mẫu gồm năm loại, hai nhà cung cấp minh họa, năm linh kiện và một phiếu nhập. Mức giá và tên nhà cung cấp trong dữ liệu này chỉ phục vụ thử nghiệm. Các màn hình ở phần sau được kết xuất trực tiếp từ thành phần Swing sau khi tải dữ liệu SQL Server.')
screens=[
('Tồn kho','Màn hình hiển thị mã, tên, loại, hãng, đơn vị, giá bán, ngưỡng và số tồn. Dữ liệu ban đầu gồm CPU01 tồn 10, RAM01 tồn 20, SSD01 tồn 12, GPU01 tồn 3 và MAIN01 tồn 8. Người dùng tìm theo mã hoặc tên, sắp xếp bảng và xuất CSV. Khi số tồn nhỏ hơn hoặc bằng ngưỡng, cột trạng thái hiển thị Cần nhập thêm.'),
('Quản lý linh kiện','Người dùng thêm hoặc sửa thông tin linh kiện, chọn loại từ danh sách có sẵn và nhập giá bán, mức tồn tối thiểu. Mã đã lưu không thay đổi. Khi linh kiện đã xuất hiện trên phiếu, khóa ngoại ngăn xóa bản ghi. Tồn kho không phải trường cho phép sửa ở form này.'),
('Quản lý loại linh kiện','Màn hình cung cấp thao tác thêm, sửa, xóa và tìm kiếm loại linh kiện. Mỗi loại có mã duy nhất và tên duy nhất. Chương trình không cho xóa loại đang được một linh kiện tham chiếu. Việc phân loại giúp dữ liệu nhất quán khi thêm các mặt hàng mới.'),
('Quản lý nhà cung cấp','Thông tin nhà cung cấp gồm mã, tên, số điện thoại và địa chỉ. Mã được chọn khi lập phiếu nhập. Khi nhà cung cấp đã có phiếu nhập, thao tác xóa bị từ chối để bảo toàn quan hệ giữa phiếu và đối tác.'),
('Danh sách phiếu nhập','Danh sách thể hiện số phiếu, ngày tạo, nhà cung cấp, người lập, ghi chú và tổng tiền. Nút Lập phiếu nhập mở form thêm nhiều dòng hàng. Người dùng cần nhập đơn giá mua thực tế; giá được điền sẵn chỉ là gợi ý từ giá bán. Nút Xem chi tiết hiển thị các linh kiện, số lượng, đơn giá và thành tiền của phiếu đã chọn.'),
('Danh sách phiếu xuất','Danh sách xuất ban đầu trống vì dữ liệu mẫu chưa có giao hàng. Khi lập phiếu, người dùng nhập người nhận và các dòng hàng. Nếu một dòng vượt tồn, toàn bộ phiếu bị từ chối. Phiếu ghi sổ thành công xuất hiện trong lịch sử, đồng thời số tồn giảm theo số lượng đã xuất.')]
for i,(title,desc) in enumerate(screens):
    page(doc);h(doc,('3.2 Giao diện chương trình' if i==0 else 'Giao diện '+title),2)
    pic(doc,OUT/f'images/screen-{i}.png',f'Hình {i+1} {title}');p(doc,desc)
page(doc);h(doc,'3.3 Một số đoạn code chính',2)
p(doc,'Đoạn mã dưới đây lấy từ Db.java. Truy vấn và tài nguyên được bao trong try-with-resources. Tham số được gán tách khỏi chuỗi SQL trước khi thực thi.')
code=(ROOT/'src/main/java/vn/edu/kholinhkien/Db.java').read_text(encoding='utf-8')
snippet=code[code.index('    public static List<Object[]> query'):code.index('    public static void execute')]
for line in snippet.splitlines():
    para=p(doc,line);para.paragraph_format.line_spacing=1;para.paragraph_format.space_after=Pt(0);para.runs[0].font.name='Consolas';para.runs[0].font.size=Pt(9)
p(doc)
p(doc,'InventoryService.post kiểm tra dữ liệu trước khi mở giao dịch. Với phiếu xuất, service khóa linh kiện rồi đọc v_TonKho. Nếu tồn ít hơn số yêu cầu, IllegalArgumentException được phát sinh; nhánh catch rollback kết nối. Sau khi các dòng hợp lệ, đầu phiếu được thêm bằng RETURN_GENERATED_KEYS để lấy mã tự tăng và ghi các chi tiết theo mã đó.')
p(doc,'Các câu lệnh cốt lõi:')
for line in ['c.setAutoCommit(false);','SELECT MaLK FROM LinhKien WITH(UPDLOCK,HOLDLOCK) WHERE MaLK=?','SELECT TonKho FROM v_TonKho WHERE MaLK=?','c.commit();','catch(Exception e) { c.rollback(); throw e; }']:
    x=p(doc,line);x.runs[0].font.name='Consolas';x.runs[0].font.size=Pt(10)
page(doc);h(doc,'3.4 Kiểm thử và đánh giá',2)
p(doc,'Bộ kiểm thử IntegrationTest chạy qua JDBC trên SQL Server thật và đạt 17 phép kiểm tra. Mỗi lần thử tạo mã riêng, sau đó dọn đúng dữ liệu đã tạo. Các dữ liệu mẫu của ứng dụng được giữ lại. Bảng dưới nhóm các kết quả theo mục tiêu kiểm tra.')
table(doc,['Nhóm kiểm thử','Kết quả'],[
('CRUD và truy vấn','Tồn ban đầu bằng 0; cập nhật tên tiếng Việt; chặn trùng khóa chính; chuỗi SQL nhập vào được xem là dữ liệu.'),
('Khóa ngoại','Chặn xóa loại và linh kiện đang được tham chiếu.'),
('Ghi sổ nhập','Nhập 10 làm tồn tăng lên 10; nhà cung cấp không tồn tại bị từ chối và tồn không đổi.'),
('Dữ liệu không hợp lệ','Từ chối số lượng 0, giá âm, dòng trùng và số tiền quá hai chữ số thập phân.'),
('Xuất vượt tồn','Yêu cầu xuất 11 khi chỉ có 10 bị chặn; không lưu đầu phiếu và tồn vẫn là 10.'),
('Xuất đồng thời','Hai yêu cầu cùng xuất 7 khi tồn 10: chỉ một yêu cầu thành công; tồn cuối bằng 3.')])
p(doc,'Kết quả cho thấy các luồng nghiệp vụ được kiểm thử giữ đúng tồn kho và quan hệ dữ liệu. Sáu màn hình cũng đã được kết xuất để kiểm tra bố cục. Bộ kiểm thử chưa bao phủ mọi sự cố môi trường như mất điện giữa lúc ghi dữ liệu hoặc mọi tương tác chuột; do đó cần thực hiện lại kịch bản nhập xuất trên máy dùng để bảo vệ đồ án.')
page(doc);h(doc,'Kết luận')
p(doc,'Đồ án đã xây dựng chương trình quản lý kho linh kiện bằng Java Swing kết nối SQL Server qua JDBC. Chương trình có bảy bảng liên kết, sáu màn hình chính, các chức năng thêm, sửa, xóa, tìm kiếm và nghiệp vụ nhập xuất có kiểm tra tồn. Cách tính tồn từ lịch sử phiếu giúp tra cứu được nguồn phát sinh số liệu.')
p(doc,'Qua quá trình xây dựng chương trình, nội dung thực hành tập trung vào thiết kế dữ liệu quan hệ, lập trình giao diện, truy vấn tham số, quản lý tài nguyên và giao dịch. Bộ kiểm thử kiểm tra được cả trường hợp bình thường lẫn vi phạm ràng buộc, trong đó có tình huống hai yêu cầu xuất đồng thời.')
p(doc,'Phạm vi hiện tại dành cho một kho và người dùng Windows được cấp quyền cơ sở dữ liệu. Chương trình chưa có tài khoản ứng dụng, phân quyền theo vai trò, serial linh kiện, bảo hành, hủy phiếu hoặc quản lý nhiều kho. Hướng phát triển là bổ sung cơ chế điều chỉnh phiếu có lịch sử, tài khoản người dùng và báo cáo nhập xuất theo khoảng thời gian. Giá trị phiếu chưa được sử dụng để tính lợi nhuận hoặc giá vốn.')
h(doc,'Tài liệu tham khảo',2)
p(doc,'1. Oracle. Java SE Desktop Overview. https://www.oracle.com/java/technologies/javase/desktop-overview.html')
p(doc,'2. Microsoft Learn. Building the connection URL with the Microsoft JDBC Driver for SQL Server. https://learn.microsoft.com/en-us/sql/connect/jdbc/building-the-connection-url')
p(doc,'3. Microsoft Learn. Download Microsoft JDBC Driver for SQL Server. https://learn.microsoft.com/en-us/sql/connect/jdbc/download-microsoft-jdbc-driver-for-sql-server')
p(doc,'4. Mẫu báo cáo đồ án Java và nội dung yêu cầu học phần do người học cung cấp. Truy cập các tài liệu trực tuyến ngày 19 tháng 10 năm 2026.')
save(doc,'Bao_cao_Java_Nguyen_Duc_Tue.docx')

db=newdoc();cover(db,'Thiết kế cơ sở dữ liệu Phần mềm Quản lý Kho Linh kiện Máy tính');page(db)
database(db,'1')
h(db,'Tạo và kiểm tra cơ sở dữ liệu',2)
p(db,'Trong SSMS, chạy sql/01_schema.sql để tạo cơ sở dữ liệu mới; chạy sql/02_seed.sql một lần nếu cần dữ liệu mẫu. SELECT * FROM v_TonKho trả về số tồn. Thiết kế đã được tạo trên SQL Server và đạt 17 phép kiểm thử JDBC. Không sửa trực tiếp bảng chi tiết trong vận hành.')
save(db,'Thiet_ke_CSDL_Kho_linh_kien.docx')
print('Created two reports; source SHA256:',original)
